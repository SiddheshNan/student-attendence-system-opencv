import cv2, sys, numpy, os
import threading
import time
from imutils.video import WebcamVideoStream
from imutils.video import FPS
import imutils
import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import mailgun

runable = True

student_chk = ['0']


def add(name, time):
    myFileName  = r'student-atten.xlsx'
    wb = load_workbook(filename=myFileName)
    ws = wb['Sheet1']
    newRowLocation = ws.max_row + 1
    ws.cell(column=1, row=newRowLocation, value=name)
    ws.cell(column=2, row=newRowLocation, value=time)
    wb.save(filename=myFileName)
    wb.close()


def stop():
    fps.stop()
    runable = False


def run():
    nam = ''
    ecount = 0
    count = 5
    detected = 0
    not_detected = 0

    # global score # make score global for this thread context
    while True:
        im = webcam.read()
        im = imutils.resize(im, width=400)
        gray = cv2.cvtColor(im, cv2.COLOR_BGR2GRAY)
        faces = face_cascade.detectMultiScale(gray, 1.3, 5)
        for (x, y, w, h) in faces:
            cv2.rectangle(im, (x, y), (x + w, y + h), (255, 0, 0), 2)
            face = gray[y:y + h, x:x + w]
            face_resize = cv2.resize(face, (width, height))
            # Try to recognize the face
            prediction = model.predict(face_resize)
            cv2.rectangle(im, (x, y), (x + w, y + h), (0, 255, 0), 3)

            if count > 0:
                if prediction[1] < 72:

                    nam = names[prediction[0]]

                    if nam in student_chk:
                        pass
                    else:
                        add(str(names[prediction[0]]),str(datetime.now().strftime("%Y-%m-%d_%H-%M-%S")))
                        student_chk.append(str(names[prediction[0]]))
                        print(nam, 'student is present')

                    count = count + 1

                    cv2.rectangle(im, (x, y), (x + w, y + h), (255, 0, 0), 3)
                    cv2.putText(im, '%s - %.0f' % (names[prediction[0]], prediction[1]), (x - 10, y - 10),
                                cv2.FONT_HERSHEY_PLAIN, 1, (0, 255, 0))

                    #print('%s - %.0f' % (names[prediction[0]], prediction[1]), (x - 10, y - 10), cv2.FONT_HERSHEY_PLAIN, 1, (0, 255, 0))
                else:
                    cv2.rectangle(im, (x, y), (x + w, y + h), (0, 0, 255), 3)
                    # print ('not recognized', (x - 10, y - 10), cv2.FONT_HERSHEY_PLAIN, 1, (255,0, 0))
                    cv2.putText(im, 'not recognized', (x - 10, y - 10), cv2.FONT_HERSHEY_PLAIN, 1, (0, 255, 0))
                    not_detected += 1
                    count -= 1

        cv2.imshow('OpenCV', im)
        key = cv2.waitKey(10)
        if key == 27:
            stop()
            break
        if key % 256 == ord('m'):
            print("sending mail.. Please wait ...")
            mailgun.mail('sid.nan23@gmail.com', 'Student Attendance System', 'Student Attendance Report at '+str(datetime.now().strftime("%Y-%m-%d_%H-%M-%S")), 'student-atten.xlsx')
            print("mail sent...")
        fps.update()





haar_file = 'haarcascade_frontalface_default.xml'
datasets = 'datasets'

print('Training...')
# Create a list of images and a list of corresponding names
(images, labels, names, id) = ([], [], {}, 0)
for (subdirs, dirs, files) in os.walk(datasets):
    for subdir in dirs:
        names[id] = subdir
        subjectpath = os.path.join(datasets, subdir)
        for filename in os.listdir(subjectpath):
            path = subjectpath + '/' + filename
            label = id
            images.append(cv2.imread(path, 0))
            labels.append(int(label))
        id += 1
(width, height) = (130, 100)

# Create a Numpy array from the two lists above
(images, labels) = [numpy.array(lis) for lis in [images, labels]]

model = cv2.face.LBPHFaceRecognizer_create()
model.train(images, labels)

# use LBPHFace recognizer on camera frame
face_cascade = cv2.CascadeClassifier(haar_file)
webcam = WebcamVideoStream(src=0).start()
fps = FPS().start()
run()
